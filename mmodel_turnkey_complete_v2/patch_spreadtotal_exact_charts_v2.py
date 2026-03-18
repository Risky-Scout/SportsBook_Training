#!/usr/bin/env python3
from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.shapes import GraphicalProperties


def build_scatter_chart(
    ws,
    title: str,
    x_title: str,
    y_title: str,
    x_col: int,
    y_col: int,
    marker_market_col: int,
    marker_model_col: int,
    min_row: int,
    max_row: int,
):
    chart = ScatterChart()
    chart.title = title
    chart.style = 2
    chart.width = 9.5
    chart.height = 4.8
    chart.x_axis.title = x_title
    chart.y_axis.title = y_title
    chart.y_axis.number_format = "0.0%"
    chart.legend.position = "r"

    xvalues = Reference(ws, min_col=x_col, min_row=min_row, max_row=max_row)
    yvalues = Reference(ws, min_col=y_col, min_row=min_row, max_row=max_row)

    dist = Series(yvalues, xvalues, title="Exact PMF")
    dist.marker.symbol = "none"
    dist.graphicalProperties = GraphicalProperties()
    dist.graphicalProperties.line.width = 19050  # ~1.5pt
    chart.series.append(dist)

    market_y = Reference(ws, min_col=marker_market_col, min_row=min_row, max_row=max_row)
    market = Series(market_y, xvalues, title="Market")
    market.graphicalProperties = GraphicalProperties()
    market.graphicalProperties.line.noFill = True
    market.marker.symbol = "diamond"
    market.marker.size = 9
    chart.series.append(market)

    model_y = Reference(ws, min_col=marker_model_col, min_row=min_row, max_row=max_row)
    model = Series(model_y, xvalues, title="Model")
    model.graphicalProperties = GraphicalProperties()
    model.graphicalProperties.line.noFill = True
    model.marker.symbol = "triangle"
    model.marker.size = 10
    chart.series.append(model)

    return chart


def main() -> None:
    ap = argparse.ArgumentParser(description="Rebuild SpreadTotal charts as trader-facing exact PMF charts.")
    ap.add_argument("--workbook", required=True)
    args = ap.parse_args()

    path = Path(args.workbook)
    wb = load_workbook(path)
    ws = wb["SpreadTotal"]

    # Remove existing chart objects and rebuild cleanly from the exact PMF tables.
    ws._charts = []

    # Margin chart from cleaned margin PMF block A:D rows 11:61
    margin_chart = build_scatter_chart(
        ws=ws,
        title="Home Margin Exact PMF | Market vs Model",
        x_title="Home margin (points)",
        y_title="Probability mass",
        x_col=1,   # A
        y_col=2,   # B
        marker_market_col=3,  # C
        marker_model_col=4,   # D
        min_row=11,
        max_row=61,
    )

    # Total chart from cleaned total PMF block F:I rows 11:61
    total_chart = build_scatter_chart(
        ws=ws,
        title="Game Total Exact PMF | Market vs Model",
        x_title="Game total (points)",
        y_title="Probability mass",
        x_col=6,   # F
        y_col=7,   # G
        marker_market_col=8,  # H
        marker_model_col=9,   # I
        min_row=11,
        max_row=61,
    )

    # Place charts below the visible ladders so they do not crowd the decision tables.
    ws.add_chart(margin_chart, "J28")
    ws.add_chart(total_chart, "J52")

    wb.save(path)
    print(f"Rebuilt SpreadTotal charts in {path}")


if __name__ == "__main__":
    main()
